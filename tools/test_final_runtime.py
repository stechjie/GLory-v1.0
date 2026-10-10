from __future__ import annotations

import copy
import difflib
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from final_runtime import (
    FILES, check_business_pages, check_code_rules, check_runtime_files,
    parse_code_rules,
)
from final_excel_owned import overlay_excel_owned, patch_json_scalars
from export_final_status import star_stat

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/balance/final status.xlsx"


class FinalRuntimeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.book = load_workbook(BOOK, read_only=True, data_only=True)
        cls.tables = {relative: json.loads((ROOT / relative).read_text(encoding="utf-8-sig"))
                      for relative in FILES}
        cls.excel_changes = overlay_excel_owned(cls.book, cls.tables)
        cls.rules = parse_code_rules(cls.book)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.book.close()

    def test_current_sources_match(self) -> None:
        self.assertEqual([], self.excel_changes)
        self.assertEqual([], check_runtime_files(ROOT, self.tables))
        self.assertEqual([], check_business_pages(self.book, self.tables))
        self.assertEqual([], check_code_rules(ROOT, self.rules))

    def test_runtime_drift_is_detected(self) -> None:
        changed = copy.deepcopy(self.tables)
        changed["data/shop.json"]["items"][0]["price"] += 1
        self.assertTrue(any("data/shop.json" in e for e in check_runtime_files(ROOT, changed)))

    def test_human_page_drift_is_detected(self) -> None:
        changed = copy.deepcopy(self.tables)
        changed["data/units/race_units.json"]["units"][0]["hp"] += 1
        self.assertTrue(any("01_棋子基础" in e for e in check_business_pages(self.book, changed)))

    def test_shop_page_drift_is_detected(self) -> None:
        changed = copy.deepcopy(self.tables)
        changed["data/shop.json"]["items"][-1]["price"] += 1
        self.assertTrue(any("93_商城目录" in e for e in check_business_pages(self.book, changed)))

    def test_avatar_frame_name_drift_is_detected(self) -> None:
        changed = copy.deepcopy(self.tables)
        changed["data/avatars.json"]["frames"][-1]["name"] = "different"
        self.assertTrue(any("shop_frame_pink_sakura frame name" in e
                            for e in check_business_pages(self.book, changed)))

    def test_cached_star_columns_match_game_scaling(self) -> None:
        units = {item["id"]: item for item in self.tables["data/units/race_units.json"]["units"]}
        for row in self.book["01_棋子基础"].iter_rows(min_row=4, values_only=True):
            if not row or not row[0]:
                continue
            unit = units[row[0]]
            for base, start, key in ((6, 6, "hp"), (10, 10, "atk"), (14, 14, "def")):
                for star in range(1, 5):
                    with self.subTest(unit=row[0], key=key, star=star):
                        self.assertEqual(star_stat(row[base], star, unit, key), row[start + star - 1])

    def test_excel_changes_only_owned_fields(self) -> None:
        book = load_workbook(BOOK, read_only=False, data_only=True)
        try:
            original = copy.deepcopy(self.tables)
            tables = copy.deepcopy(original)
            unit = tables["data/units/race_units.json"]["units"][0]
            original_skill = copy.deepcopy(unit["star4"])
            unit_id = unit["id"]
            for row in book["01_棋子基础"].iter_rows(min_row=4):
                if row[0].value == unit_id:
                    row[6].value += 10
                    break
            changes = overlay_excel_owned(book, tables)
            updated = tables["data/units/race_units.json"]["units"][0]
            self.assertEqual(original["data/units/race_units.json"]["units"][0]["hp"] + 10, updated["hp"])
            self.assertEqual(original_skill, updated["star4"])
            self.assertEqual(1, len(changes))
            self.assertEqual([], check_runtime_files(ROOT, original))
            self.assertTrue(check_runtime_files(ROOT, tables))
            self.assertEqual(3 * updated["hp"], star_stat(updated["hp"], 3, updated, "hp"))
            source = (ROOT / "data/units/race_units.json").read_text(encoding="utf-8-sig")
            patched = patch_json_scalars(source, changes, tables["data/units/race_units.json"])
            diff = list(difflib.unified_diff(source.splitlines(), patched.splitlines()))
            self.assertEqual(2, len([line for line in diff if line.startswith(("+", "-")) and not line.startswith(("+++", "---"))]))
            self.assertEqual(tables["data/units/race_units.json"], json.loads(patched))
        finally:
            book.close()

    def test_invalid_excel_input_stops_before_writing(self) -> None:
        book = load_workbook(BOOK, read_only=False, data_only=True)
        try:
            book["05_野怪"]["E4"] = -1
            with self.assertRaisesRegex(ValueError, "positive whole number"):
                overlay_excel_owned(book, copy.deepcopy(self.tables))
        finally:
            book.close()

    def test_boss_and_pet_excel_values_override_json(self) -> None:
        book = load_workbook(BOOK, read_only=False, data_only=True)
        try:
            boss = book["06_Boss"]
            boss["D4"] = 3450  # 2300 base HP after the 1.5 battle multiplier
            pet = book["03_宠物与商城"]
            pet["E4"] = "12%"
            pet["F4"] = str(pet["F4"].value).replace("10%", "12%")
            pet["N4"] = str(pet["N4"].value).replace("10%", "12%")
            tables = copy.deepcopy(self.tables)
            changes = overlay_excel_owned(book, tables)
            self.assertEqual(2300, tables["data/boss/bosses.json"]["bosses"][0]["hp"])
            self.assertEqual(0.12, tables["data/pets/pets.json"]["pets"][0]["value"])
            self.assertEqual(2, len(changes))
            self.assertEqual(240, tables["data/boss/bosses.json"]["bosses"][0]["skill_damage"])
        finally:
            book.close()

    def test_code_rule_drift_is_detected(self) -> None:
        changed = copy.deepcopy(self.rules)
        changed["pet_draw.price"]["value"] += 1
        self.assertTrue(any("pet_draw.price" in e for e in check_code_rules(ROOT, changed)))


if __name__ == "__main__":
    unittest.main()
