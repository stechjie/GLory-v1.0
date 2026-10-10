"""Export the presentation catalog from docs/balance/final status.xlsx.

Usage: python tools/export_final_status.py [--apply-runtime | --check]
Requires openpyxl. The workbook is the editable source; the JSON is packaged by Godot.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook

from final_runtime import (FILES, check_business_pages, check_code_rules,
                           check_runtime_files, parse_code_rules)
from final_excel_owned import overlay_excel_owned, patch_json_scalars

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/balance/final status.xlsx"
OUTPUT = ROOT / "data/balance/final_status.json"


def records(ws):
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row and row[0] not in (None, ""):
            yield list(row)


def text(value):
    return "" if value is None else str(value)


def star_stat(base, star, unit, key):
    multiplier = (1.0, 1.5, 3.0, 3.0 * max(1.0, float(unit.get("star4_multiplier", 1.15))))[star - 1]
    value = max(0 if key == "def" else 1, math.floor(base * multiplier + 0.5))
    return 1 if key == "atk" and unit["id"] == "human_death_servant" else value


def build():
    wb = load_workbook(BOOK, read_only=True, data_only=True)
    tables = {relative: json.loads((ROOT / relative).read_text(encoding="utf-8-sig"))
              for relative in FILES}
    changes = overlay_excel_owned(wb, tables)
    rules = parse_code_rules(wb)
    code_errors = check_code_rules(ROOT, rules)
    if code_errors:
        raise ValueError("final workbook code rules disagree with source:\n" + "\n".join(code_errors))
    page_errors = check_business_pages(wb, tables)
    if page_errors:
        raise ValueError("final workbook pages disagree with runtime JSON:\n" + "\n".join(page_errors[:20]))
    source_cell = text(wb["00_版本与口径"]["C5"].value)
    runtime_cell = text(wb["00_版本与口径"]["B6"].value)
    source_hash = re.search(r"[0-9a-f]{64}", source_cell)
    runtime_stamp = re.search(r"20\d\d-\d\d-\d\dT\d\d:\d\d:\d\d", runtime_cell)
    if not source_hash or not runtime_stamp:
        raise ValueError("Workbook provenance is missing from 00_版本与口径")
    out = {
        "schema_version": 1,
        "title": "final status",
        "source_workbook": text(wb["00_版本与口径"]["B5"].value),
        "source_sha256": source_hash.group(),
        "source_runtime_utc": runtime_stamp.group(),
        "scope": "本机 Godot 运行值与用户指定文案目标；线上服务端逐项系数待核",
        "units": {}, "pets": {}, "synergies": {}, "treasures": {},
        "linkages": {}, "shop_items": {},
    }
    basics = {text(r[0]): r for r in records(wb["01_棋子基础"])}
    unit_defs = {item["id"]: item for item in tables["data/units/race_units.json"]["units"]}
    for r in records(wb["02_棋子技能"]):
        unit_id = text(r[0])
        b = basics[unit_id]
        unit = unit_defs[unit_id]
        out["units"][unit_id] = {
            "name_cn": text(r[1]), "race": text(b[2]), "element": text(b[3]),
            "tier": int(b[4]), "cost": int(b[5]), "skill_id": text(r[3]),
            "skill_name_cn": text(r[2]), "skill_cn_1to3": text(r[6]),
            "skill_cn_4": text(r[12]), "skill_en_1to3": text(r[11]),
            "skill_raw_params": text(r[7]), "star4_changes": text(r[8]),
            "stats": {
                str(star): {
                    "hp": star_stat(int(b[6]), star, unit, "hp"),
                    "atk": star_stat(int(b[10]), star, unit, "atk"),
                    "defense": star_stat(int(b[14]), star, unit, "def"),
                }
                for star in range(1, 5)
            },
        }
        if unit_id == "dark_dragon":
            out["units"][unit_id]["skill_damage_multiplier_1to3"] = unit["damage_atk_pct"]
            out["units"][unit_id]["skill_damage_multiplier_4"] = unit["star4"]["damage_atk_pct"]
    for r in records(wb["03_宠物与商城"]):
        pet_id = text(r[0])
        out["pets"][pet_id] = {
            "name_cn": text(r[1]), "name_en": text(r[2]),
            "effect_type": text(r[3]), "value_percent": text(r[4]),
            "detail_cn": text(r[5]), "scope_cn": text(r[6]),
            "starter": text(r[7]) == "是", "summary_cn": text(r[13]),
            "shop_item_id": text(r[9]), "shop_currency": text(r[10]),
            "shop_price": r[11], "shop_enabled": text(r[12]) == "是",
        }
    for r in records(wb["04_羁绊"]):
        if text(r[4]) != "已审定":
            continue
        out["synergies"].setdefault(text(r[0]), []).append({
            "threshold": int(r[1]), "name_cn": text(r[2]),
            "detail_cn": text(r[3]),
        })
    for r in records(wb["09_宝藏与套装"]):
        if text(r[7]) != "已审定":
            continue
        out["treasures"][text(r[0])] = {
            "name_cn": text(r[1]), "category_cn": text(r[2]),
            "effect_cn": text(r[4]), "hu_pai_text_cn": text(r[5]),
        }
    for r in records(wb["10_宝藏联动"]):
        if text(r[8]) != "已审定":
            continue
        out["linkages"][text(r[0])] = {
            "name_cn": text(r[1]), "requires_cn": text(r[2]),
            "effect_cn": text(r[4]),
        }
    for item in tables["data/shop.json"]["items"]:
        out["shop_items"][item["id"]] = {
            "kind": item["kind"], "grants": item["grants"],
            "currency": item["currency"], "price": item["price"],
            "enabled": item.get("enabled", True),
            "name_cn": item.get("name", ""), "name_en": item.get("name_en", ""),
        }
    expected = {"units": 40, "pets": 5, "synergies": 5,
                "treasures": 25, "linkages": 11}
    for key, count in expected.items():
        if len(out[key]) != count:
            raise ValueError(f"{key}: expected {count}, got {len(out[key])}")
    payload = json.dumps({"display": out, "runtime": tables, "code_rules": rules}, ensure_ascii=False,
                         sort_keys=True, separators=(",", ":")).encode("utf-8")
    out["balance_version"] = hashlib.sha256(payload).hexdigest()
    out["final_workbook_sha256"] = hashlib.sha256(BOOK.read_bytes()).hexdigest()
    return out, tables, changes


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true",
                        help="Fail when JSON differs from the workbook")
    parser.add_argument("--apply-runtime", action="store_true",
                        help="Write approved runtime JSON tables from the workbook")
    args = parser.parse_args()
    if args.check and args.apply_runtime:
        parser.error("--check and --apply-runtime cannot be combined")
    try:
        generated, tables, changes = build()
    except (ValueError, KeyError, OSError, json.JSONDecodeError) as exc:
        parser.exit(2, f"FINAL_STATUS_ERROR: {exc}\n")
    if args.check:
        current = json.loads(OUTPUT.read_text(encoding="utf-8"))
        runtime_errors = check_runtime_files(ROOT, tables)
        if current != generated or runtime_errors:
            details = ["final_status.json differs from final status.xlsx"] if current != generated else []
            details.extend(runtime_errors[:20])
            raise SystemExit("\n".join(details))
        print("FINAL_STATUS_EXPORT_CHECK_OK")
    else:
        if args.apply_runtime:
            pending = {}
            for relative in FILES:
                file_changes = [item for item in changes if item["file"] == relative]
                if file_changes:
                    path = ROOT / relative
                    pending[path] = patch_json_scalars(path.read_text(encoding="utf-8-sig"),
                                                       file_changes, tables[relative])
            for change in changes:
                print(f"Excel -> JSON {change['sheet']} {change['id']}.{change['field']}: "
                      f"{change['old']} -> {change['new']}")
            for path, source in pending.items():
                path.write_text(source, encoding="utf-8", newline="")
                print(f"Updated runtime {path.relative_to(ROOT)}")
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        current = json.loads(OUTPUT.read_text(encoding="utf-8")) if OUTPUT.exists() else None
        if current != generated:
            OUTPUT.write_text(json.dumps(generated, ensure_ascii=False, indent=2) + "\n",
                              encoding="utf-8", newline="\n")
            print(f"Exported {OUTPUT}")
        else:
            print("Final status catalog already current")


if __name__ == "__main__":
    main()
